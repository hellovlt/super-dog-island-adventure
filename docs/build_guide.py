"""Rebuild the English Word guide from the maintained player guide."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
language = OxmlElement('w:lang')
language.set(qn('w:val'), 'en-US')
normal.element.get_or_add_rPr().append(language)
doc.core_properties.title = 'Super Dog 3D - Game Guide'
doc.core_properties.subject = 'English game description, controls, and local access'
doc.core_properties.author = ''

def text_runs(paragraph, text):
    for part in re.split(r'(https?://[^\s]+)', text):
        if part.startswith('http'):
            url = part.rstrip('.')
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            style = OxmlElement('w:rStyle'); style.set(qn('w:val'), 'Hyperlink')
            props.append(style); run.append(props)
            value = OxmlElement('w:t'); value.text = url
            run.append(value); link.append(run); paragraph._p.append(link)
            if part.endswith('.'): paragraph.add_run('.')
        else:
            paragraph.add_run(part)

for line in (root / 'Super-Dog-3D-Guide.md').read_text(encoding='utf-8').splitlines():
    line = line.strip()
    if not line:
        continue
    line = line.replace('**', '').replace('`', '')
    if line.startswith('# '):
        doc.add_heading(line[2:], 0)
    elif line.startswith('## '):
        doc.add_heading(line[3:], 1)
    elif line.startswith('- '):
        text_runs(doc.add_paragraph(style='List Bullet'), line[2:])
    elif re.match(r'^\d+\. ', line):
        text_runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', line))
    else:
        text_runs(doc.add_paragraph(), line)

output = root / 'Super Dog Adventure Guide.docx'
doc.save(output)
print(output)
